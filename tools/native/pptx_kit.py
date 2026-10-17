"""Helpers for building *native* PowerPoint slides in the house style.

Native slides are for content that is mostly text, cards and lists building in, where real
PowerPoint shapes are simpler, editable and crisp. Remotion clips stay for slides whose motion carries
meaning (see the skill's house-style.md, "Native PowerPoint or Remotion?").

Coordinates are 1920×1080 "pixels", the same as the Remotion scenes: on a 13.333×7.5 in slide,
1 px = 6350 EMU and 1 px of font size = 0.5 pt, so layouts can be copied across directly.
Each builder records entrance animations with `k.anim(shape_name, at_seconds, effect)`;
tools/animate_deck.ps1 applies them through PowerPoint, timed from the start of the slide.
"""

from __future__ import annotations

import os

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

FONT = "Segoe UI"  # Inter (used in the Remotion clips) is not installed on the build machine

MONO = "Cascadia Mono"

C = {
    "bg": "0A0F1E", "panel": "111A30", "panelEdge": "24314F", "text": "F2F5FB", "muted": "8F9BB8",
    "faint": "4A5675", "tmf": "E6007E", "component": "F5A623", "canvas": "2F6FED", "operator": "1FC7A8",
    "api": "8B6CFF", "identity": "FF6B6B", "observability": "38BDF8", "secrets": "E9C46A",
    "events": "A3E635", "ai": "F472B6", "carbon": "34D399", "cost": "FB923C", "agent": "F8FAFC",
    "ok": "22C55E", "danger": "F43F5E",
}

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def px(v: float) -> Emu:
    return Emu(int(round(v * 6350)))


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6.lstrip("#").upper())


def _alpha(fill_or_line_color_parent, alpha: float) -> None:
    """Add <a:alpha> to the srgbClr under a fill (0..1 opacity)."""
    clr = fill_or_line_color_parent.find(f".//{{{A}}}srgbClr")
    if clr is not None and alpha < 1:
        el = etree.SubElement(clr, f"{{{A}}}alpha")
        el.set("val", str(int(alpha * 100000)))


class Kit:
    def __init__(self, slide, assets_dir: str):
        self.slide = slide
        self.shapes = slide.shapes
        self.assets = assets_dir
        self.anims: list[dict] = []
        self._n = 0

    # ---------- animation plan ----------
    def anim(self, name: str, at: float, effect: str = "fade", dur: float = 0.6) -> None:
        """effect: fade | rise (fly in from bottom) | wipe-up | zoom (faded zoom)."""
        self.anims.append({"shape": name, "at": round(at, 2), "effect": effect, "dur": dur})

    def _name(self, shape, name: str | None):
        self._n += 1
        shape.name = name or f"Shape {self._n}"
        return shape

    # ---------- primitives ----------
    def background(self, file: str = "stage.png"):
        pic = self.shapes.add_picture(os.path.join(self.assets, file), 0, 0, px(1920), px(1080))
        self._name(pic, "Background")
        # send to back
        tree = self.shapes._spTree
        tree.remove(pic._element)
        tree.insert(2, pic._element)
        return pic

    def image(self, name, file, x, y, w, h=None):
        pic = self.shapes.add_picture(os.path.join(self.assets, file), px(x), px(y), px(w), px(h) if h else None)
        return self._name(pic, name)

    def rect(self, name, x, y, w, h, fill=None, fill_alpha=1.0, line=None, line_w=3, radius=18, line_alpha=1.0):
        shape = self.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
        if radius:
            shape.adjustments[0] = min(0.5, radius / min(w, h))
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(C.get(fill, fill))
            _alpha(shape.fill._xPr.find(f"{{{A}}}solidFill"), fill_alpha)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = _rgb(C.get(line, line))
            shape.line.width = px(line_w)
            _alpha(shape.line._get_or_add_ln().find(f"{{{A}}}solidFill"), line_alpha)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return self._name(shape, name)

    def circle(self, name, cx, cy, d, fill=None, fill_alpha=1.0, line=None, line_w=4):
        shape = self.shapes.add_shape(MSO_SHAPE.OVAL, px(cx - d / 2), px(cy - d / 2), px(d), px(d))
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(C.get(fill, fill))
            _alpha(shape.fill._xPr.find(f"{{{A}}}solidFill"), fill_alpha)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = _rgb(C.get(line, line))
            shape.line.width = px(line_w)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return self._name(shape, name)

    def text(self, name, x, y, w, h, runs, size=32, color="text", bold=False, align="left",
             anchor="top", spacing=0, line_spacing=1.15, font=None, shape=None):
        """runs: a string, or a list of paragraphs; each paragraph a string or a list of (text, {opts})."""
        if shape is None:
            shape = self.shapes.add_textbox(px(x), px(y), px(w), px(h))
            self._name(shape, name)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        paras = runs if isinstance(runs, list) else [runs]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = line_spacing
            parts = para if isinstance(para, list) else [(para, {})]
            for txt, opt in parts:
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name = opt.get("font", font or FONT)
                f.bold = opt.get("bold", bold)
                f.size = Pt(opt.get("size", size) * 0.5)
                f.color.rgb = _rgb(C.get(opt.get("color", color), opt.get("color", color)))
                sp = opt.get("spacing", spacing)
                if sp:
                    r._r.get_or_add_rPr().set("spc", str(int(sp * 50)))  # px → hundredths of a point
        return shape

    @staticmethod
    def pill_width(label: str, size: int = 26) -> int:
        """Estimated pill width for Segoe UI bold (used to lay pills out in a row)."""
        return int(len(label) * size * 0.56 + 44)

    def pill(self, name, x, y, label, color, size=26, w=None, text_color="text", link=None):
        width = w or self.pill_width(label, size)
        h = int(size * 1.75)
        shape = self.rect(name, x, y, width, h, fill=color, fill_alpha=0.15, line=color, line_w=2, radius=h / 2)
        self.text(None, 0, 0, 0, 0, label, size=size, color=text_color, bold=True, align="center", anchor="middle", shape=shape)
        if link:
            shape.click_action.hyperlink.address = link
        return shape

    def line(self, name, x1, y1, x2, y2, color="muted", width=4, arrow=False, dash=False, elbow=False):
        """Straight (default) or elbow connector; elbow routes horizontal → vertical → horizontal."""
        kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
        conn = self.shapes.add_connector(kind, px(x1), px(y1), px(x2), px(y2))
        conn.line.color.rgb = _rgb(C.get(color, color))
        conn.line.width = px(width)
        ln = conn.line._get_or_add_ln()
        if dash:
            d = etree.SubElement(ln, f"{{{A}}}prstDash")
            d.set("val", "dash")
        if arrow:
            t = etree.SubElement(ln, f"{{{A}}}tailEnd")
            t.set("type", "triangle")
            t.set("w", "med")
            t.set("h", "med")
        return self._name(conn, name)

    def group(self, name, shapes):
        g = self.shapes.add_group_shape(shapes)
        return self._name(g, name)

    def heading(self, kicker: str, title: str, color: str = "tmf", at: float = 0.0):
        k = self.text("Kicker", 120, 84, 1500, 40, kicker.upper(), size=26, color=color, bold=True, spacing=5)
        t = self.text("Title", 120, 122, 1700, 90, title, size=68, bold=True)
        self.anim("Kicker", at, "fade", 0.5)
        self.anim("Title", at + 0.1, "rise", 0.6)
        return k, t
