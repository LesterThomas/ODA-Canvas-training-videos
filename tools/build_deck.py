"""Assemble a training-video deck: each slide is either a full-bleed auto-playing Remotion clip or a
native PowerPoint slide, and every slide carries its speaker notes.

Usage:
    python tools/build_deck.py <notes.json> <clips_dir> <posters_dir> <clip_prefix> <output.pptx> [title]
                               [--native tools/native/<deck>.py --assets videos/out/assets]

- --native     module with BUILDERS = {slide number: builder(kit)}; those slides are drawn as native
               PowerPoint shapes (tools/native/pptx_kit.py) instead of embedding a clip
- --assets     PNG assets for native slides (node videos/scripts/assets.mjs)
Writes <output>.anim.json next to the deck: the native slides' entrance animations and every slide's
narration length. Then run tools/animate_deck.ps1 to apply them in PowerPoint.

- notes.json   list of {"n": slide number, "title": ..., "notes": ...} (made from scripts/<deck-id>.md)
- clips_dir    rendered Remotion clips named <clip_prefix>-NN-*.mp4
- posters_dir  full-resolution final-frame PNGs named <clip_prefix>-NN-*.png (used as the poster
               frame, so a static export of the slide shows the finished diagram)
"""

import argparse
import glob
import importlib.util
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}

# Auto-play the movie when the slide starts (python-pptx's default timing is click-to-play).
AUTOPLAY_TIMING = """<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
    <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
      <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
        <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
          <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">
            <p:stCondLst><p:cond delay="0"/></p:stCondLst>
            <p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{dur}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst>
          </p:cTn></p:par>
        </p:childTnLst></p:cTn></p:par>
      </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn>
    <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
    <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
    </p:seq>
    <p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
  </p:childTnLst></p:cTn></p:par></p:tnLst>
</p:timing>"""


def one(pattern: str) -> str:
    hits = sorted(glob.glob(pattern))
    if len(hits) != 1:
        raise SystemExit(f"expected exactly one match for {pattern}, got {hits}")
    return hits[0]


def mp4_duration_ms(path: str) -> int:
    """Read the movie duration from the mp4 'mvhd' box (no ffprobe needed)."""
    with open(path, "rb") as f:
        data = f.read()
    i = data.find(b"mvhd")
    version = data[i + 4]
    if version == 1:
        timescale = int.from_bytes(data[i + 24 : i + 28], "big")
        duration = int.from_bytes(data[i + 28 : i + 36], "big")
    else:
        timescale = int.from_bytes(data[i + 16 : i + 20], "big")
        duration = int.from_bytes(data[i + 20 : i + 24], "big")
    return round(duration * 1000 / timescale)


def load_builders(path: str | None) -> dict:
    if not path:
        return {}
    sys.path.insert(0, os.path.dirname(os.path.abspath(path)))  # so the module can import pptx_kit
    spec = importlib.util.spec_from_file_location("native_builders", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.BUILDERS


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("notes_json")
    ap.add_argument("clips_dir")
    ap.add_argument("posters_dir")
    ap.add_argument("prefix")
    ap.add_argument("out_path")
    ap.add_argument("title", nargs="?", default="")
    ap.add_argument("--native", default=None)
    ap.add_argument("--assets", default="videos/out/assets")
    a = ap.parse_args()
    slides = json.load(open(a.notes_json, encoding="utf-8"))
    builders = load_builders(a.native)
    if builders:
        from pptx_kit import Kit  # noqa: E402 (path added by load_builders)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    if a.title:
        prs.core_properties.title = a.title
    layout = prs.slide_layouts[5]  # "Title Only": keeps a real title for outline/accessibility
    plan = {"deck": os.path.basename(a.out_path), "slides": {}}

    for idx, s in enumerate(slides, 1):
        n = s["n"]
        slide = prs.slides.add_slide(layout)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(0x0A, 0x0F, 0x1E)

        # Title lives just off the slide: present in outline view and for screen readers,
        # never rendered on top of the content.
        title = slide.shapes.title
        title.text = s["title"]
        title.left, title.top = prs.slide_width + Inches(0.5), Inches(0)
        title.width, title.height = Inches(6), Inches(1)

        entry = {"n": n, "secs": s.get("secs"), "kind": "video"}
        if n in builders:
            kit = Kit(slide, a.assets)
            builders[n](kit)
            entry.update(kind="native", anims=kit.anims)
        else:
            clip = one(os.path.join(a.clips_dir, f"{a.prefix}-{n:02d}-*.mp4"))
            poster = one(os.path.join(a.posters_dir, f"{a.prefix}-{n:02d}-*.png"))
            movie = slide.shapes.add_movie(
                clip, 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=poster, mime_type="video/mp4"
            )
            movie.name = f"Clip {n:02d}"
            # Replace click-to-play timing with auto-play.
            sld = slide.element
            for t in sld.findall("p:timing", NS):
                sld.remove(t)
            timing = etree.fromstring(AUTOPLAY_TIMING.format(dur=mp4_duration_ms(clip), spid=movie.shape_id))
            ext = sld.find("p:extLst", NS)  # schema order: cSld, clrMapOvr, transition, timing, extLst
            if ext is not None:
                ext.addprevious(timing)
            else:
                sld.append(timing)
        plan["slides"][str(idx)] = entry
        slide.notes_slide.notes_text_frame.text = s["notes"]

    lock = os.path.join(os.path.dirname(os.path.abspath(a.out_path)), "~$" + os.path.basename(a.out_path))
    try:
        if os.path.exists(lock):
            raise PermissionError
        prs.save(a.out_path)
    except PermissionError:
        sys.exit(f"{a.out_path} is open in PowerPoint: close it there, then build again")
    plan_path = os.path.splitext(a.out_path)[0] + ".anim.json"
    json.dump(plan, open(plan_path, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    native = sum(1 for e in plan["slides"].values() if e["kind"] == "native")
    print(f"saved {a.out_path} with {len(slides)} slides ({native} native, {len(slides) - native} video); plan {plan_path}")


if __name__ == "__main__":
    main()
