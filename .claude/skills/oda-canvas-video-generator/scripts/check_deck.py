"""Structural check of a built training-video deck against its notes JSON.

Usage: python check_deck.py decks/<deck-id>.pptx decks/<deck-id>.notes.json

Checks each slide has notes identical to the JSON. Video slides need exactly one media shape and the
auto-play timing (mainSeq + playFrom) that build_deck.py writes. Native slides (listed in
<deck>.anim.json) need no media, and every planned animation must name a shape on the slide. Run
after tools/animate_deck.ps1 to also confirm the animations were applied. Exits non-zero on any problem.
"""

import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def main() -> None:
    deck, notes_json = sys.argv[1], sys.argv[2]
    prs = Presentation(deck)
    expected = json.load(open(notes_json, encoding="utf-8"))
    plan_path = os.path.splitext(deck)[0] + ".anim.json"
    plan = json.load(open(plan_path, encoding="utf-8"))["slides"] if os.path.exists(plan_path) else {}
    problems = []
    if len(prs.slides) != len(expected):
        problems.append(f"deck has {len(prs.slides)} slides, notes JSON has {len(expected)}")
    for i, (slide, exp) in enumerate(zip(prs.slides, expected), 1):
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes != exp["notes"]:
            problems.append(f"slide {i}: notes differ from script")
        media = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.MEDIA]
        timing = slide.element.find(f"{P}timing")
        xml = etree.tostring(timing).decode() if timing is not None else ""
        entry = plan.get(str(i), {})
        if entry.get("kind") == "native":
            if media:
                problems.append(f"slide {i}: native slide has {len(media)} media shapes")
            names = {s.name for s in slide.shapes}
            missing = [a["shape"] for a in entry.get("anims", []) if a["shape"] not in names]
            if missing:
                problems.append(f"slide {i}: animated shapes not found: {missing}")
            if entry.get("anims") and 'nodeType="mainSeq"' not in xml:
                problems.append(f"slide {i}: animations not applied yet (run tools/animate_deck.ps1)")
            continue
        if len(media) != 1:
            problems.append(f"slide {i}: {len(media)} media shapes (expected 1)")
        if 'nodeType="mainSeq"' not in xml or "playFrom" not in xml:
            problems.append(f"slide {i}: no auto-play timing")
    for p in problems:
        print("PROBLEM:", p)
    print(f"{len(prs.slides)} slides checked, {len(problems)} problems")
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
